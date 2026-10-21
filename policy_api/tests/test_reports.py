import asyncio
import io
import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.demo import run_demo
from app.excel import build_workbook
from app.reports import compose
from app.reports.charts import (change_matrix_png, custom_chart_png, impact_types_png, status_by_policy_png,
                                validate_chart_spec)
from app.reports.compare import build_comparison, compare_policies
from app.reports.docx_report import build_docx
from app.reports.pdf_report import build_pdf
from app.reports.redline import redline

PNG = b"\x89PNG\r\n\x1a\n"
BAD_TITLE = '<font color="red">x</font><a href="http://e">y</a>'


@pytest.fixture(scope="module")
def result(tmp_path_factory) -> dict:
    samples = Path(__file__).resolve().parents[2] / "samples"
    res, _ = asyncio.run(run_demo(samples, tmp_path_factory.mktemp("demo") / "d.xlsx"))
    return json.loads(json.dumps(res))


def pdf_text(data: bytes) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)


def docx_text(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])


SPEC = {"title": "Training plan", "subtitle": "Q4", "sections": [
    {"heading": "Overview", "paragraphs": ["Plain paragraph."], "bullets": ["one", "two"]},
    {"heading": "Numbers", "table": {"columns": ["Name", "Count"], "rows": [["A", 1], ["B", 2.5]]},
     "chart": {"type": "bar", "title": "Counts", "labels": ["A", "B"], "series": [{"name": "n", "values": [1, 2.5]}]}},
]}


def test_redline_roundtrip():
    parts = redline("the quick brown fox", "the slow brown dog")
    assert "".join(t for op, t in parts if op != "insert") == "the quick brown fox"
    assert "".join(t for op, t in parts if op != "delete") == "the slow brown dog"
    assert ("delete", "quick") in parts and ("insert", "slow") in parts


def test_charts_are_png(result):
    for fn in (status_by_policy_png, change_matrix_png, impact_types_png):
        assert fn(result).startswith(PNG)
    empty = {**result, "rows": [], "changes": [], "counts": {**result["counts"], "not_impacted": 0}}
    for fn in (status_by_policy_png, change_matrix_png, impact_types_png):
        assert fn(empty).startswith(PNG)


@pytest.mark.parametrize("kind", ["bar", "stacked_bar", "line", "pie"])
def test_custom_chart_kinds(kind):
    spec = {"type": kind, "title": "T", "labels": ["a", "b", "c"], "series": [{"name": "s", "values": [1, 2, 3]}],
            "x_label": "x", "y_label": "y"}
    assert custom_chart_png(spec).startswith(PNG)


@pytest.mark.parametrize("mutate", [
    lambda s: s.update(type="radar"),
    lambda s: s.update(labels=["a"] * 51, series=[{"name": "s", "values": [1] * 51}]),
    lambda s: s.update(series=[{"name": "s", "values": [1, 2]}]),
    lambda s: s.update(series=[{"name": str(i), "values": [1, 2, 3]} for i in range(11)]),
    lambda s: s.update(series=[{"name": "s", "values": [1, float("nan"), 3]}]),
    lambda s: s.update(series=[{"name": "s", "values": [1, "2", 3]}]),
    lambda s: s.update(title="x" * 500),
    lambda s: s.pop("title"),
])
def test_custom_chart_rejects_bad_input(mutate):
    spec = {"type": "bar", "title": "T", "labels": ["a", "b", "c"], "series": [{"name": "s", "values": [1, 2, 3]}]}
    mutate(spec)
    with pytest.raises(ValueError):
        validate_chart_spec(spec)
    with pytest.raises(ValueError):
        custom_chart_png(spec)


def test_pdf_from_demo(result):
    text = pdf_text(build_pdf(result))
    assert "SAF-112" in text and "What changed" in text and "How this was checked" in text
    assert "Page 1 of" in text and result["run_id"] in text


def test_pdf_shows_markup_literally(result):
    bad = json.loads(json.dumps(result))
    bad["rows"][0]["policy_title"] = BAD_TITLE
    bad["rows"][0]["why"] = "<b>bold</b> & <a href='x'>link</a>"
    bad["changes"][0]["new_text"] = "new <font color='red'>text</font> § “quoted” → …"
    text = "".join(pdf_text(build_pdf(bad)).split())  # PDF text wraps lines; compare without spaces
    assert "".join(BAD_TITLE.split()) in text
    assert "<b>bold</b>" in text and "color='red'>text</font>" in text and "“quoted”" in text


def test_docx_from_demo(result):
    data = build_docx(result)
    d = Document(io.BytesIO(data))
    assert "SAF-112" in docx_text(data)
    assert len(d.inline_shapes) == 3
    assert d.tables and d.tables[1].rows[0].cells[0].text == "Policy"
    runs = [r for p in d.paragraphs for r in p.runs]
    assert any(r.font.strike for r in runs) and any(r.font.underline for r in runs)


def test_xlsx_has_new_sheets(result):
    wb = load_workbook(io.BytesIO(build_workbook(result)))
    assert {"Summary", "Impacted", "Charts", "Change matrix", "Method"} <= set(wb.sheetnames)
    assert len(wb["Charts"]._charts) == 1
    m = wb["Change matrix"]
    values = {c.value for row in m.iter_rows(min_row=2, min_col=2) for c in row}
    assert "Yes: needs updating" in values and "Maybe: check by hand" in values


def test_xlsx_formula_in_policy_title_stays_text(result):
    bad = json.loads(json.dumps(result))
    bad["rows"][0]["policy_title"] = '=HYPERLINK("http://e","x")'
    wb = load_workbook(io.BytesIO(build_workbook(bad)))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f"


@pytest.mark.parametrize("fmt,check", [("pdf", lambda b: b.startswith(b"%PDF")), ("docx", lambda b: b[:2] == b"PK"),
                                       ("xlsx", lambda b: b[:2] == b"PK"), ("png", lambda b: b.startswith(PNG))])
def test_custom_all_formats(fmt, check):
    assert check(compose.build_custom(SPEC, fmt))


def test_custom_carries_disclaimer():
    assert "Prepared by the assistant at your request" in pdf_text(compose.build_custom(SPEC, "pdf")).replace("\n", " ")
    assert "Prepared by the assistant at your request" in docx_text(compose.build_custom(SPEC, "docx"))
    ws = load_workbook(io.BytesIO(compose.build_custom(SPEC, "xlsx")))["Document"]
    assert any("Prepared by the assistant at your request" in str(c.value) for row in ws.iter_rows() for c in row)
    assert len(ws._images) == 1


def test_custom_pdf_markup_is_literal():
    spec = {"title": BAD_TITLE, "sections": [{"heading": "<u>h</u>", "paragraphs": ["<b>p</b>"]}]}
    text = "".join(pdf_text(compose.build_custom(spec, "pdf")).split())
    assert "".join(BAD_TITLE.split()) in text and "<u>h</u>" in text and "<b>p</b>" in text


def test_custom_xlsx_formula_is_text():
    spec = {"title": "=HYPERLINK(\"http://e\",\"t\")", "sections": [
        {"heading": "=1+1", "paragraphs": ["=HYPERLINK(\"http://evil\",\"click\")"], "bullets": ["=SUM(A1)"],
         "table": {"columns": ["=A"], "rows": [["=HYPERLINK(\"http://e\",\"x\")"], [3]]}}]}
    ws = load_workbook(io.BytesIO(compose.build_custom(spec, "xlsx")))["Document"]
    found = 0
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                found += 1
                assert c.data_type == "s", c.coordinate
    assert found >= 5


def _spec(**sec):
    return {"title": "T", "sections": [sec or {"paragraphs": ["x"]}]}


def test_custom_limits():
    with pytest.raises(ValueError, match="too many sections"):
        compose.build_custom({"title": "T", "sections": [{"paragraphs": ["x"]}] * 41}, "pdf")
    with pytest.raises(ValueError, match="too many rows"):
        compose.build_custom(_spec(table={"columns": ["a"], "rows": [["x"]] * 2001}), "xlsx")
    with pytest.raises(ValueError, match="too many columns"):
        compose.build_custom(_spec(table={"columns": ["a"] * 31, "rows": []}), "xlsx")
    with pytest.raises(ValueError, match="too long"):
        compose.build_custom(_spec(paragraphs=["x" * 5001]), "pdf")
    with pytest.raises(ValueError, match="too much text"):
        compose.build_custom({"title": "T", "sections": [{"paragraphs": ["x" * 4000] * 20} for _ in range(5)]}, "pdf")
    with pytest.raises(ValueError):
        compose.build_custom(SPEC, "exe")
    with pytest.raises(ValueError, match="no chart"):
        compose.build_custom(_spec(paragraphs=["x"]), "png")
    with pytest.raises(ValueError):
        compose.build_custom(_spec(table={"columns": ["a", "b"], "rows": [["x"]]}), "pdf")


def test_compare_finds_pol007_changes(samples):
    old = next((samples / "policies").glob("POL-007*.md")).read_text()
    new = (samples / "changes" / "POL-007-v4.md").read_text()
    cmp = compare_policies(old, new)
    by = {s["number"]: s for s in cmp["sections"]}
    assert by["2.2"]["status"] == "changed"
    assert ("delete", "24") in by["2.2"]["redline"] and ("insert", "12") in by["2.2"]["redline"]
    assert by["4.2"]["status"] == "changed" and by["5.3"]["status"] == "added"
    assert cmp["counts"]["same"] >= 1
    assert "2.2" in {n for n, s in by.items() if s["status"] != "same"}
    assert "24" in pdf_text(build_comparison("pdf", "v3", old, "v4", new))
    assert "Authorized worker" in docx_text(build_comparison("docx", "v3", old, "v4", new))
    wb = load_workbook(io.BytesIO(build_comparison("xlsx", "v3", old, "v4", new)))
    assert wb.sheetnames == ["Summary", "Sections"]
    statuses = [r[2].value for r in wb["Sections"].iter_rows(min_row=2)]
    assert "Changed" in statuses and "Added" in statuses and "Same" in statuses


def test_compare_removed_and_empty():
    cmp = compare_policies("## 1 Alpha\none two three.\n## 2 Beta\nfour five six.", "## 1 Alpha\none two three.")
    assert {s["number"]: s["status"] for s in cmp["sections"]} == {"1": "same", "2": "removed"}
    with pytest.raises(ValueError):
        compare_policies("", "x")


def test_dollar_signs_in_chart_text_are_plain_text():
    # Unbalanced or nested "$" would raise in matplotlib's mathtext parser if it were on.
    spec = {"type": "bar", "title": "Costs $\\frac{$ and $x^{2", "labels": ["$a", "b$$\\sqrt{"],
            "series": [{"name": "$$", "values": [1, 2]}]}
    assert custom_chart_png(spec).startswith(b"\x89PNG")


def test_total_table_cells_are_capped():
    table = {"columns": ["n"] * 30, "rows": [[1] * 30] * 2000}
    spec = {"title": "Big", "sections": [{"heading": f"S{i}", "table": table} for i in range(2)]}
    with pytest.raises(ValueError, match="too many cells"):
        compose.validate_spec(spec)


def test_compare_of_thousands_of_sections_finishes_quickly():
    import time
    from app.reports.compare import compare_policies
    left = "".join(f"1.{i} Rule {i} for area {i * 31 % 997}.\n" for i in range(1, 4000))
    right = "".join(f"2.{i} Rule {i} for area {i * 31 % 997}.\n" for i in range(1, 4000))
    start = time.monotonic()
    cmp = compare_policies(left, right)
    assert time.monotonic() - start < 5
    assert cmp["counts"]["removed"] == 3999 and cmp["counts"]["added"] == 3999
