import asyncio
import io

from openpyxl import load_workbook

from app.demo import run_demo
from app.excel import build_workbook


def _status(result: dict) -> dict[str, str]:
    out: dict[str, str] = {}
    for r in result["rows"]:
        if out.get(r["policy_id"]) != "Yes":
            out[r["policy_id"]] = r["status"]
    return out


def test_demo_finds_the_expected_policies(samples, tmp_path):
    result, answer = asyncio.run(run_demo(samples, tmp_path / "out.xlsx"))
    assert _status(result) == {"SAF-112": "Yes", "SAF-130": "Yes", "SAF-150": "Maybe"}
    assert {p["policy_id"] for p in result["not_impacted"]} == {"SAF-141", "SAF-160"}
    c = result["counts"]
    assert (c["subset_policies"], c["impacted"], c["maybe"], c["not_impacted"], c["changes"]) == (5, 2, 1, 2, 4)
    assert "HR-201" not in answer and "POL-007" in answer
    assert answer.startswith("**2 of the 5 Safety policies need updating")


def test_weekly_check_is_flagged_with_evidence(samples, tmp_path):
    result, _ = asyncio.run(run_demo(samples, tmp_path / "out.xlsx"))
    row = next(r for r in result["rows"] if r["policy_id"] == "SAF-112" and r["clause"] == "2.2")
    assert row["status"] == "Yes"
    assert "weekly" in row["evidence_candidate"]
    assert "§4.2" in row["head_clauses"]


def test_workbook_sheets_and_counts(samples, tmp_path):
    asyncio.run(run_demo(samples, tmp_path / "out.xlsx"))
    wb = load_workbook(tmp_path / "out.xlsx")
    assert wb.sheetnames == ["Summary", "Impacted", "Head policy changes", "Checked, not impacted", "Method", "Charts", "Change matrix"]
    assert wb["Impacted"]["A2"].value == "Yes"
    assert wb["Head policy changes"].max_row == 5  # header + 4 changes
    assert "DEMO DATA" in str([c.value for c in wb["Summary"]["A"]])


def test_policy_text_never_becomes_a_formula(samples, tmp_path):
    result, _ = asyncio.run(run_demo(samples, tmp_path / "out.xlsx"))
    result["rows"][0]["clause_text"] = '=HYPERLINK("http://evil.example","click")'
    wb = load_workbook(io.BytesIO(build_workbook(result)))
    cell = wb["Impacted"]["J2"]
    assert cell.value.startswith("=HYPERLINK")
    assert cell.data_type == "s"
