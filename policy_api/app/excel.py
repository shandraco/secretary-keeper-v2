"""The Excel comparison workbook. Built entirely in code from the run result; the model never writes cells."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

HEADER = PatternFill("solid", fgColor="1F3A5F")
FILL = {"Yes": PatternFill("solid", fgColor="F8D7DA"), "Maybe": PatternFill("solid", fgColor="FFF3CD")}
WRAP = Alignment(wrap_text=True, vertical="top")
IMPACT_LABEL = {
    "conflict": "Conflicts with the new wording",
    "outdated_reference": "Points to a moved or removed section",
    "changed_term": "Uses a term whose meaning changed",
    "needs_update": "Repeats old wording",
    "redundant": "Now repeats the head policy",
    "none": "",
}
CHANGE_LABEL = {"modified": "Changed", "added": "Added", "removed": "Removed", "renumbered": "Renumbered"}


def _neutralise_formulas(ws: Worksheet) -> None:
    """Policy text starting with '=' must stay text, never become a live formula."""
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                cell.data_type = "s"


def _table(ws: Worksheet, headers: list[str], rows: list[list], widths: list[int]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = WRAP
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def _summary(ws: Worksheet, result: dict) -> None:
    c, head = result["counts"], result["head"]
    ws.title = "Summary"
    ws["A1"] = "Policy impact check"
    ws["A1"].font = Font(bold=True, size=16)
    if c["impacted"] or c["maybe"]:
        answer = (f"{c['impacted']} of {c['subset_policies']} policies need updating. "
                  f"{c['maybe']} might be affected and need a person to check. {c['not_impacted']} are not affected.")
    else:
        answer = f"None of the {c['subset_policies']} policies are affected."
    items = [
        ("Question", f"Which {result['subset']} policies are affected by the change to {head['id']} ({head['title']})?"),
        ("Answer", answer),
        ("Compared", f"{head['old_label']} → {head['new_label']}"),
        ("Need updating", c["impacted"]),
        ("Might be affected (check by hand)", c["maybe"]),
        ("Not affected", c["not_impacted"]),
        ("Changed sections in the head policy", c["changes"]),
        ("Section pairs checked", c["pairs_checked"]),
        ("Could not be checked automatically", c["errors"]),
        ("Created", result["finished_at"]),
        ("Run ID", result["run_id"]),
        ("How to read this file",
         "'Impacted' lists every affected section. Red rows need updating; yellow rows need a person to decide. "
         "Fill in 'Review decision' as you go. 'Checked, not impacted' shows what was ruled out."),
    ]
    for i, (label, value) in enumerate(items, start=3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value).alignment = WRAP
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 100
    if result["method"].get("demo_mode"):
        ws.cell(row=len(items) + 4, column=1, value="DEMO DATA: produced by stand-in models, not real ones.").font = Font(bold=True, color="C00000")


STATUS_WORD = {"Yes": "Yes: needs updating", "Maybe": "Maybe: check by hand"}


def _charts_sheet(wb: Workbook, result: dict) -> None:
    """Native Excel charts, so staff can click on them and see the numbers behind."""
    ws = wb.create_sheet("Charts")
    per: dict[str, list] = {}
    for r in result["rows"]:
        p = per.setdefault(r["policy_id"], [0, 0])
        if r["status"] == "Yes":
            p[0] += 1
        elif r["status"] == "Maybe":
            p[1] += 1
    ids = sorted(per, key=lambda k: (-per[k][0], -per[k][1], k))
    ws.append(["Policy", "Needs updating (sections)", "Might be affected (sections)"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = WRAP
    for k in ids:
        ws.append([k, per[k][0], per[k][1]])
    ws.append([f"{result['counts']['not_impacted']} not affected", 0, 0])
    n = len(ids) + 1
    for col, w in zip("ABC", (26, 26, 28)):
        ws.column_dimensions[col].width = w
    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "Affected sections, by policy"
    chart.y_axis.title = "Number of sections"
    chart.y_axis.majorUnit = 1
    chart.y_axis.delete = False
    chart.x_axis.delete = False
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=n + 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
    chart.series[0].graphicalProperties.solidFill = "B42318"
    chart.series[1].graphicalProperties.solidFill = "B7791F"
    chart.legend.position = "b"
    chart.height, chart.width = max(6, 1.2 * n + 3), 18
    ws.add_chart(chart, "E2")


def _matrix_sheet(wb: Workbook, result: dict) -> None:
    ws = wb.create_sheet("Change matrix")
    labels = [c["label"] for c in result["changes"]]
    pols: dict[str, str] = {}
    for r in result["rows"]:
        pols.setdefault(r["policy_id"], r["policy_title"])
    cols = sorted(pols)
    rank = {"Maybe": 1, "Yes": 2}
    grid: dict[tuple[str, str], str] = {}
    for r in result["rows"]:
        if r["status"] in rank and r["head_change"] in labels:
            key = (r["head_change"], r["policy_id"])
            if rank[r["status"]] > rank.get(grid.get(key, ""), 0):
                grid[key] = r["status"]
    ws.append(["Change in the head policy"] + [f"{p}\n{pols[p]}" for p in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")
    for label in labels:
        ws.append([label] + [STATUS_WORD.get(grid.get((label, p)), "") for p in cols])
    ws.column_dimensions["A"].width = 34
    for i in range(2, len(cols) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 24
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = WRAP
        row[0].font = Font(bold=True)
        for cell in row[1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            fill = FILL.get((cell.value or "").split(":")[0])
            if fill:
                cell.fill = fill
    ws.freeze_panes = "B2"
    if not labels or not cols:
        ws["A2"] = "No changes affected any policy."


def build_workbook(result: dict) -> bytes:
    wb = Workbook()
    _summary(wb.active, result)

    ws = wb.create_sheet("Impacted")
    _table(
        ws,
        ["Status", "Policy", "Policy title", "Owner", "Section", "Section title", "Problem", "Head policy section(s)",
         "What changed in the head policy (new wording)", "What this section says", "Why it's affected",
         "Evidence (head policy)", "Evidence (this policy)", "Suggested edit", "How it was found",
         "Checked by larger model", "Review decision", "Reviewer"],
        [[r["status"], r["policy_id"], r["policy_title"], r["owner"], f"§{r['clause']}", r["clause_heading"],
          IMPACT_LABEL.get(r["impact_type"], ""), r["head_clauses"], r["head_new_text"] or "(section removed)",
          r["clause_text"], r["why"], r["evidence_head"], r["evidence_candidate"], r["suggested_edit"],
          r["found_by"], "Yes" if r["escalated"] else "No", "", ""] for r in result["rows"]],
        [9, 11, 24, 18, 9, 20, 24, 14, 50, 50, 40, 28, 28, 40, 22, 12, 16, 16],
    )
    for row in ws.iter_rows(min_row=2):
        fill = FILL.get(row[0].value)
        if fill:
            for cell in row:
                cell.fill = fill
    if result["rows"]:
        dv = DataValidation(type="list", formula1='"Agree,Disagree,Not sure"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"Q2:Q{len(result['rows']) + 1}")

    _table(
        wb.create_sheet("Head policy changes"),
        ["Section", "Change", "Section title", "Old wording", "New wording", "Defined terms in this section"],
        [[f"§{c['number']}" + (f" (was §{c['old_number']})" if c.get("old_number") else ""),
          CHANGE_LABEL.get(c["kind"], c["kind"]), c["heading"], c["old_text"], c["new_text"], ", ".join(c["terms"])]
         for c in result["changes"]],
        [16, 12, 24, 60, 60, 28],
    )
    _table(
        wb.create_sheet("Checked, not impacted"),
        ["Policy", "Policy title", "Owner", "Reason"],
        [[p["policy_id"], p["title"], p["owner"], p["reason"]] for p in result["not_impacted"]],
        [12, 36, 20, 60],
    )
    _table(
        wb.create_sheet("Method"),
        ["Setting", "Value"],
        [[k, str(v)] for k, v in result["method"].items()] + [["run_id", result["run_id"]]],
        [24, 60],
    )

    _charts_sheet(wb, result)
    _matrix_sheet(wb, result)

    for sheet in wb.worksheets:
        _neutralise_formulas(sheet)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
