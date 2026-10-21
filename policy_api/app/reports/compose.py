"""Documents the chat assistant writes on request. The spec is data; nothing in it is ever read as markup."""
from __future__ import annotations

import io
import math

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..excel import _neutralise_formulas
from . import ACCENT, DISCLAIMER, clean
from .charts import custom_chart_png, validate_chart_spec

MAX_SECTIONS, MAX_ROWS, MAX_COLS = 40, 2000, 30
MAX_STRING, MAX_TOTAL = 5000, 300_000
MAX_CELLS = 20_000  # across all tables in one document
FORMATS = ("pdf", "docx", "xlsx", "png")


class _Budget:
    def __init__(self) -> None:
        self.total = 0
        self.cells = 0

    def text(self, value, what: str, required: bool = False) -> str:
        if value is None or value == "":
            if required:
                raise ValueError(f"The document needs a {what}.")
            return ""
        if not isinstance(value, str):
            raise ValueError(f"The {what} must be text.")
        if len(value) > MAX_STRING:
            raise ValueError(f"The {what} is too long (the limit is {MAX_STRING} characters for each piece of text).")
        self.total += len(value)
        if self.total > MAX_TOTAL:
            raise ValueError("The document has too much text (the limit is about 300,000 characters).")
        return clean(value)


def _list(value, what: str) -> list:
    if value is None:
        return []
    if not isinstance(value, list):
        raise ValueError(f"The {what} must be a list.")
    return value


def validate_spec(spec) -> dict:
    """Strict check. Returns a cleaned copy or raises ValueError with a plain-English message."""
    if not isinstance(spec, dict):
        raise ValueError("The document description is not in the expected form.")
    b = _Budget()
    out = {"title": b.text(spec.get("title"), "title", True), "subtitle": b.text(spec.get("subtitle"), "subtitle"), "sections": []}
    sections = _list(spec.get("sections"), "sections")
    if not sections:
        raise ValueError("The document needs at least one section.")
    if len(sections) > MAX_SECTIONS:
        raise ValueError(f"The document has too many sections (the limit is {MAX_SECTIONS}).")
    for s in sections:
        if not isinstance(s, dict):
            raise ValueError("Each section must have a heading, paragraphs, bullets, a table or a chart.")
        sec = {"heading": b.text(s.get("heading"), "section heading"),
               "paragraphs": [b.text(p, "paragraph") for p in _list(s.get("paragraphs"), "paragraphs")],
               "bullets": [b.text(p, "bullet point") for p in _list(s.get("bullets"), "bullets")],
               "table": None, "chart": None}
        t = s.get("table")
        if t is not None:
            if not isinstance(t, dict):
                raise ValueError("A table needs columns and rows.")
            cols = _list(t.get("columns"), "table columns")
            rows = _list(t.get("rows"), "table rows")
            if not cols:
                raise ValueError("A table needs at least one column.")
            if len(cols) > MAX_COLS:
                raise ValueError(f"A table has too many columns (the limit is {MAX_COLS}).")
            if len(rows) > MAX_ROWS:
                raise ValueError(f"A table has too many rows (the limit is {MAX_ROWS}).")
            b.cells += len(rows) * len(cols)
            if b.cells > MAX_CELLS:
                raise ValueError(f"The document's tables have too many cells (the limit is {MAX_CELLS:,} in total).")
            clean_rows = []
            for r in rows:
                if not isinstance(r, list) or len(r) != len(cols):
                    raise ValueError("Every table row must have one value for each column.")
                row = []
                for v in r:
                    if isinstance(v, bool) or v is None:
                        raise ValueError("Table cells must be text or numbers.")
                    if isinstance(v, (int, float)):
                        if not math.isfinite(v):
                            raise ValueError("Table numbers must be finite.")
                        row.append(v)
                    else:
                        row.append(b.text(v, "table cell") if isinstance(v, str) else _bad())
                clean_rows.append(row)
            sec["table"] = {"columns": [b.text(c, "column heading") for c in cols], "rows": clean_rows}
        if s.get("chart") is not None:
            ch = validate_chart_spec(s["chart"])
            b.total += len(ch["title"]) + sum(len(x) for x in ch["labels"])
            if b.total > MAX_TOTAL:
                raise ValueError("The document has too much text (the limit is about 300,000 characters).")
            sec["chart"] = ch
        out["sections"].append(sec)
    return out


def _bad():
    raise ValueError("Table cells must be text or numbers.")


def _fmt_cell(v) -> str:
    return f"{v:,.10g}" if isinstance(v, float) else (f"{v:,}" if isinstance(v, int) else v)


def _pdf(spec: dict) -> bytes:
    from reportlab.platypus import KeepTogether, Paragraph, Spacer
    from . import pdf_report as P
    st = P.styles()
    story = [Paragraph(P.esc(spec["title"]), st["title"])]
    if spec["subtitle"]:
        story.append(Paragraph(P.esc(spec["subtitle"]), st["muted"]))
    story += [Spacer(1, 4), Paragraph(P.esc(DISCLAIMER), st["small"]), Spacer(1, 6)]
    for s in spec["sections"]:
        if s["heading"]:
            story.append(Paragraph(P.esc(s["heading"]), st["h1"]))
        for p in s["paragraphs"]:
            story += [Paragraph(P.esc(p), st["body"]), Spacer(1, 4)]
        if s["bullets"]:
            story.append(P.bullets(st, s["bullets"]))
            story.append(Spacer(1, 4))
        if s["table"]:
            t = s["table"]
            n = len(t["columns"])
            story.append(P.table(st, t["columns"], [[_fmt_cell(v) for v in r] for r in t["rows"]], [1] * n))
            story.append(Spacer(1, 6))
        if s["chart"]:
            story.append(KeepTogether([P.image(custom_chart_png(s["chart"]), P.CONTENT_W * 0.9)]))
    return P.render(story, spec["title"], clean(spec["title"], 60) + "  ·  prepared by the assistant")


def _docx(spec: dict) -> bytes:
    from docx.shared import Pt
    from . import docx_report as D
    doc = D.new_document(spec["title"], clean(spec["title"], 60) + "  ·  prepared by the assistant")
    doc.add_paragraph(spec["title"], style="Title")
    if spec["subtitle"]:
        doc.add_paragraph(spec["subtitle"])
    doc.add_paragraph().add_run(DISCLAIMER).font.size = Pt(9)
    for s in spec["sections"]:
        if s["heading"]:
            doc.add_heading(s["heading"], level=1)
        for p in s["paragraphs"]:
            doc.add_paragraph(p)
        for b in s["bullets"]:
            doc.add_paragraph(b, style="List Bullet")
        if s["table"]:
            t = s["table"]
            D.add_table(doc, t["columns"], [[_fmt_cell(v) for v in r] for r in t["rows"]], [1] * len(t["columns"]))
        if s["chart"]:
            D.add_picture(doc, custom_chart_png(s["chart"]), 6.0)
    return D.to_bytes(doc)


def _xlsx(spec: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Document"
    wrap = Alignment(wrap_text=True, vertical="top")
    ws["A1"] = spec["title"]
    ws["A1"].font = Font(bold=True, size=16)
    row = 2
    if spec["subtitle"]:
        ws.cell(row=row, column=1, value=spec["subtitle"])
        row += 1
    ws.cell(row=row, column=1, value=DISCLAIMER).font = Font(italic=True, color="5B6670")
    row += 2
    ncols = max([len(s["table"]["columns"]) for s in spec["sections"] if s["table"]] + [1])
    for s in spec["sections"]:
        if s["heading"]:
            ws.cell(row=row, column=1, value=s["heading"]).font = Font(bold=True, size=13, color=ACCENT)
            row += 1
        for text in s["paragraphs"] + ["•  " + b for b in s["bullets"]]:
            c = ws.cell(row=row, column=1, value=text)
            c.alignment = wrap
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(ncols, 4))
            ws.row_dimensions[row].height = min(400, 15 * (1 + len(text) // 110))
            row += 1
        if s["table"]:
            t = s["table"]
            for j, name in enumerate(t["columns"], 1):
                c = ws.cell(row=row, column=j, value=name)
                c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=ACCENT), wrap
            row += 1
            for r in t["rows"]:
                for j, v in enumerate(r, 1):
                    ws.cell(row=row, column=j, value=v).alignment = wrap
                row += 1
            row += 1
        if s["chart"]:
            img = XlImage(io.BytesIO(custom_chart_png(s["chart"])))
            img.width, img.height = img.width * 0.5, img.height * 0.5
            ws.add_image(img, f"A{row}")
            row += int(img.height / 20) + 2
        row += 1
    for j in range(1, max(ncols, 4) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 28 if ncols > 1 else 100
    _neutralise_formulas(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _png(spec: dict) -> bytes:
    for s in spec["sections"]:
        if s["chart"]:
            return custom_chart_png(s["chart"], DISCLAIMER)
    raise ValueError("There is no chart in this document to save as a picture.")


def build_custom(spec: dict, fmt: str) -> bytes:
    if fmt not in FORMATS:
        raise ValueError("The document can be made as a PDF, Word, Excel or PNG file.")
    clean_spec = validate_spec(spec)
    return {"pdf": _pdf, "docx": _docx, "xlsx": _xlsx, "png": _png}[fmt](clean_spec)
