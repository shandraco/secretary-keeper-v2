"""Word report (python-docx). Plain text only: python-docx writes runs as text, so no markup is interpreted."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import (ACCENT, DELETE_RGB, INSERT_RGB, STATUS_FILL, STATUS_STRONG, STATUS_WORD, clean, headline)
from .charts import change_matrix_png, impact_types_png, status_by_policy_png
from .pdf_report import _date
from .redline import redline


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def new_document(title: str, footer: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    doc.core_properties.title = clean(title, 200)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 11.5)):
        s = doc.styles[name]
        s.font.name, s.font.size, s.font.bold = "Calibri", Pt(size), True
        s.font.color.rgb = _rgb(ACCENT) if name != "Heading 2" else _rgb("1F2933")
        rf = s.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(attr)) is not None:
                del rf.attrib[qn(attr)]
    p = sec.footer.paragraphs[0]
    r = p.add_run(clean(footer, 160) + "   ·   Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = _rgb("5B6670")
    _field(p, "PAGE")
    return doc


def _field(paragraph, code: str) -> None:
    run = paragraph.add_run()
    run.font.size = Pt(8)
    run.font.color.rgb = _rgb("5B6670")
    for kind, text in (("begin", None), (None, code), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_table(doc, header: list[str], rows: list[list], widths: list[float], status_col: int | None = None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(clean(h))
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = _rgb("FFFFFF")
        shade(cell, ACCENT)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trPr.append(hdr)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(clean(v))
            run.font.size = Pt(9.5)
            if status_col == i:
                key = {STATUS_WORD["Yes"]: "Yes", STATUS_WORD["Maybe"]: "Maybe"}.get(str(v))
                if key:
                    shade(cells[i], STATUS_FILL[key])
    total = sum(widths)
    usable = 6.7
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(usable * w / total)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(usable * w / total)
    doc.add_paragraph()
    return t


def add_redline(doc, old: str, new: str):
    para = doc.add_paragraph()
    prev = ""
    for op, text in redline(clean(old), clean(new)):
        if op == "insert" and prev == "delete":
            para.add_run(" ")
        run = para.add_run(text)
        if op == "delete":
            run.font.strike = True
            run.font.color.rgb = _rgb(DELETE_RGB)
        elif op == "insert":
            run.font.underline = True
            run.font.color.rgb = _rgb(INSERT_RGB)
        prev = op
    return para


def add_legend(doc) -> None:
    p = doc.add_paragraph()
    r = p.add_run("How to read the wording: ")
    r.font.size = Pt(9)
    a = p.add_run("red crossed-out text")
    a.font.strike, a.font.size, a.font.color.rgb = True, Pt(9), _rgb(DELETE_RGB)
    p.add_run(" was removed; ").font.size = Pt(9)
    b = p.add_run("blue underlined text")
    b.font.underline, b.font.size, b.font.color.rgb = True, Pt(9), _rgb(INSERT_RGB)
    p.add_run(" is new.").font.size = Pt(9)


def add_picture(doc, png: bytes, width_in: float = 6.2) -> None:
    doc.add_picture(io.BytesIO(png), width=Inches(width_in))


def to_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_docx(result: dict) -> bytes:
    c, head, m = result["counts"], result["head"], result["method"]
    doc = new_document(f"Policy impact check {head['id']}", f"Run {result['run_id']}  ·  {_date(result)}")
    doc.add_paragraph("Policy impact check", style="Title")
    p = doc.add_paragraph(clean(f"{head['id']} {head['title']}: {head['old_label']} → {head['new_label']}"))
    p.runs[0].font.color.rgb = _rgb("5B6670")
    p.runs[0].font.size = Pt(9.5)
    lead = doc.add_paragraph()
    lr = lead.add_run(headline(result))
    lr.font.size = Pt(13)
    if m.get("demo_mode"):
        d = doc.add_paragraph().add_run("DEMO DATA: produced by stand-in models, not real ones. Never use for real decisions.")
        d.bold = True
        d.font.color.rgb = _rgb(STATUS_STRONG["Yes"])

    add_table(doc, ["Need updating", "Might be affected", "Not affected", "Changed sections"],
              [[c["impacted"], c["maybe"], c["not_impacted"], c["changes"]]], [1, 1, 1, 1])
    if c["errors"]:
        doc.add_paragraph(f"{c['errors']} section pairs could not be checked automatically and are marked “Might be affected”.")

    doc.add_heading("At a glance", level=1)
    add_picture(doc, status_by_policy_png(result), 5.8)
    add_picture(doc, change_matrix_png(result), 5.4)
    if any(r["status"] in ("Yes", "Maybe") for r in result["rows"]):
        add_picture(doc, impact_types_png(result), 5.6)

    yes = [r for r in result["rows"] if r["status"] == "Yes"]
    doc.add_heading("Needs updating", level=1)
    if yes:
        add_table(doc, ["Policy", "Section", "Status", "Why"],
                  [[f"{r['policy_id']}\n{r['policy_title']}", f"§{r['clause']}", STATUS_WORD["Yes"], r["why"]] for r in yes],
                  [23, 11, 14, 52], status_col=2)
    else:
        doc.add_paragraph("No section needs updating.")

    doc.add_heading("What changed in the head policy", level=1)
    add_legend(doc)
    if result["changes"]:
        for ch in result["changes"]:
            doc.add_heading(clean(ch["label"]), level=2).paragraph_format.keep_with_next = True
            add_redline(doc, ch["old_text"], ch["new_text"])
    else:
        doc.add_paragraph("No wording changed.")

    maybe = [r for r in result["rows"] if r["status"] == "Maybe"]
    doc.add_heading("Might be affected", level=1)
    if maybe:
        doc.add_paragraph("A person should read these and decide.")
        add_table(doc, ["Policy", "Section", "Status", "Why"],
                  [[f"{r['policy_id']}\n{r['policy_title']}", f"§{r['clause']}", STATUS_WORD["Maybe"], r["why"]] for r in maybe],
                  [23, 11, 14, 52], status_col=2)
    else:
        doc.add_paragraph("Nothing needs a second look.")

    doc.add_heading("Checked, not affected", level=1)
    if result["not_impacted"]:
        add_table(doc, ["Policy", "Owner", "Reason"],
                  [[f"{x['policy_id']}\n{x['title']}", x.get("owner") or "", x["reason"]] for x in result["not_impacted"]], [28, 20, 52])
    else:
        doc.add_paragraph("Every policy in the group had at least one section to look at.")

    doc.add_heading("How this was checked", level=1)
    for line in (
        f"Each changed section of {head['id']} was compared with sections in the other {result['subset']} policies. {c['pairs_checked']} section pairs were looked at.",
        f"Sections were found by what they cite and what they mean (matching: {m.get('embedder')}, {m.get('reranker')}).",
        f"A checking model ({m.get('judge_model')}) judged each pair; {c['escalated']} unclear pairs were passed to a second model ({m.get('escalate_model')}).",
        "Documents, charts and tables are built by code from these results. The model does not write them.",
        "A person should still read every “Needs updating” and “Might be affected” item before changing a policy.",
    ):
        doc.add_paragraph(clean(line), style="List Bullet")
    return to_bytes(doc)
