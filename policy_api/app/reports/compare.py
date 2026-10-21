"""Side-by-side comparison of two versions of one policy (or two policies), as PDF, Word or Excel."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..diff import _sort_key, diff_clauses
from ..documents import split_clauses, split_front_matter
from ..excel import _neutralise_formulas, _table
from . import DELETE_RGB, INSERT_RGB, STATUS_FILL, clean
from .redline import redline

MAX_POLICY_CHARS = 400_000
STATUS_WORD = {"same": "Same", "changed": "Changed", "added": "Added", "removed": "Removed", "renumbered": "Renumbered"}
STATUS_TINT = {"changed": STATUS_FILL["Maybe"], "added": "DDEFE3", "removed": STATUS_FILL["Yes"], "renumbered": "E6ECF5"}


def compare_policies(left_text: str, right_text: str) -> dict:
    """Align sections by number (and by similarity when renumbered). Returns {"counts", "sections"}."""
    for name, t in (("first", left_text), ("second", right_text)):
        if not isinstance(t, str) or not t.strip():
            raise ValueError(f"The {name} policy has no text to compare.")
        if len(t) > MAX_POLICY_CHARS:
            raise ValueError(f"The {name} policy is too long to compare here.")
    old = split_clauses(split_front_matter(clean(left_text))[1])
    new = split_clauses(split_front_matter(clean(right_text))[1])
    changes = diff_clauses(old, new)
    old_by, new_by = {c.number: c for c in old}, {c.number: c for c in new}
    sections: list[dict] = []
    changed_new = {c.number for c in changes if c.kind != "removed"}
    for ch in changes:
        status = "changed" if ch.kind == "modified" else ch.kind
        sections.append({"number": ch.number, "old_number": ch.old_number, "heading": ch.heading, "status": status,
                         "left": ch.old_text, "right": ch.new_text})
    for c in new:
        if c.number not in changed_new and c.number in old_by:
            sections.append({"number": c.number, "old_number": None, "heading": c.heading, "status": "same",
                             "left": old_by[c.number].text, "right": c.text})
    sections.sort(key=lambda s: _sort_key(s["number"]))
    for s in sections:
        s["redline"] = redline(s["left"], s["right"]) if s["status"] in ("changed", "renumbered") else None
    counts = {k: sum(1 for s in sections if s["status"] == k) for k in STATUS_WORD}
    counts["total"] = len(sections)
    return {"counts": counts, "sections": sections}


def _name(s: dict) -> str:
    if s["number"] == "0":
        return "Introduction"
    n = f"§{s['number']}"
    if s["old_number"]:
        n += f" (was §{s['old_number']})"
    return n + (f" {s['heading']}" if s["heading"] else "")


def _summary_sentence(cmp: dict, left: str, right: str) -> str:
    c = cmp["counts"]
    if not (c["changed"] or c["added"] or c["removed"] or c["renumbered"]):
        return f"{left} and {right} have the same wording in every section."
    parts = [f"{c[k]} {w}" for k, w in (("changed", "changed"), ("added", "added"), ("removed", "removed"), ("renumbered", "renumbered")) if c[k]]
    return f"Compared with {left}, {right} has " + ", ".join(parts) + f". {c['same']} of {c['total']} sections are unchanged."


def _pdf(cmp, ll, rl) -> bytes:
    from reportlab.platypus import KeepTogether, Paragraph, Spacer
    from . import pdf_report as P
    st = P.styles()
    story = [Paragraph("Policy comparison", st["title"]), Paragraph(P.esc(f"{ll} → {rl}"), st["muted"]), Spacer(1, 10),
             Paragraph(P.esc(_summary_sentence(cmp, ll, rl)), st["lead"]), P.legend_para(st), Spacer(1, 6)]
    shown = [s for s in cmp["sections"] if s["status"] != "same"]
    if not shown:
        story.append(Paragraph("Nothing differs.", st["body"]))
    for s in shown:
        title = Paragraph(P.esc(f"{_name(s)}: {STATUS_WORD[s['status']]}"), st["h2"])
        if s["redline"] is not None:
            body = Paragraph(P.redline_markup(s["left"], s["right"]), st["body"])
        elif s["status"] == "added":
            body = Paragraph(f'<font color="#{INSERT_RGB}"><u>{P.esc(s["right"])}</u></font>', st["body"])
        else:
            body = Paragraph(f'<font color="#{DELETE_RGB}"><strike>{P.esc(s["left"])}</strike></font>', st["body"])
        story.append(KeepTogether([title, body]))
    same = [s for s in cmp["sections"] if s["status"] == "same"]
    if same:
        story.append(Paragraph("Unchanged sections", st["h1"]))
        story.append(Paragraph(P.esc(", ".join(_name(s) for s in same)), st["body"]))
    return P.render(story, "Policy comparison", clean(f"{ll} → {rl}", 120))


def _docx(cmp, ll, rl) -> bytes:
    from docx.shared import Pt
    from . import docx_report as D
    doc = D.new_document("Policy comparison", clean(f"{ll} → {rl}", 120))
    doc.add_paragraph("Policy comparison", style="Title")
    doc.add_paragraph(clean(f"{ll} → {rl}"))
    doc.add_paragraph().add_run(_summary_sentence(cmp, ll, rl)).font.size = Pt(13)
    D.add_legend(doc)
    shown = [s for s in cmp["sections"] if s["status"] != "same"]
    for s in shown:
        doc.add_heading(clean(f"{_name(s)}: {STATUS_WORD[s['status']]}"), level=2).paragraph_format.keep_with_next = True
        if s["redline"] is not None:
            D.add_redline(doc, s["left"], s["right"])
        elif s["status"] == "added":
            D.add_redline(doc, "", s["right"])
        else:
            D.add_redline(doc, s["left"], "")
    same = [s for s in cmp["sections"] if s["status"] == "same"]
    if same:
        doc.add_heading("Unchanged sections", level=1)
        doc.add_paragraph(clean(", ".join(_name(s) for s in same)))
    return D.to_bytes(doc)


def _xlsx(cmp, ll, rl) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Policy comparison"
    ws["A1"].font = Font(bold=True, size=16)
    rows = [("Compared", f"{clean(ll)} → {clean(rl)}"), ("Answer", _summary_sentence(cmp, clean(ll), clean(rl)))]
    rows += [(STATUS_WORD[k], cmp["counts"][k]) for k in STATUS_WORD] + [("Total sections", cmp["counts"]["total"])]
    for i, (a, b) in enumerate(rows, 3):
        ws.cell(row=i, column=1, value=a).font = Font(bold=True)
        ws.cell(row=i, column=2, value=b)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 100
    sec = wb.create_sheet("Sections")
    _table(sec, ["Section", "Title", "Status", clean(ll, 60), clean(rl, 60)],
           [[f"§{s['number']}" + (f" (was §{s['old_number']})" if s["old_number"] else ""), clean(s["heading"]),
             STATUS_WORD[s["status"]], clean(s["left"]), clean(s["right"])] for s in cmp["sections"]],
           [14, 24, 12, 70, 70])
    for row in sec.iter_rows(min_row=2):
        tint = STATUS_TINT.get(next((k for k, w in STATUS_WORD.items() if w == row[2].value), ""))
        if tint:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=tint)
    for sheet in wb.worksheets:
        _neutralise_formulas(sheet)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_comparison(fmt: str, left_label: str, left_text: str, right_label: str, right_text: str) -> bytes:
    if fmt not in ("pdf", "docx", "xlsx"):
        raise ValueError("The comparison can be made as a PDF, Word or Excel file.")
    ll, rl = clean(left_label, 80) or "First version", clean(right_label, 80) or "Second version"
    cmp = compare_policies(left_text, right_text)
    return {"pdf": _pdf, "docx": _docx, "xlsx": _xlsx}[fmt](cmp, ll, rl)
